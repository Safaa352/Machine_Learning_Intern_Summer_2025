from pptx import Presentation
from pptx.util import Inches
import matplotlib.pyplot as plt
import seaborn as sns
import pandas as pd
import numpy as np
import os

from sklearn.preprocessing import LabelEncoder
from sklearn.model_selection import train_test_split
from sklearn.neighbors import KNeighborsClassifier
from sklearn.metrics import accuracy_score, classification_report, confusion_matrix, ConfusionMatrixDisplay

# ----------------------------
# Step 1: Load and preprocess
# ----------------------------
df = pd.read_csv("first inten project.csv")
df['is_canceled'] = df['booking status'].apply(lambda x: 1 if x.strip() == 'Canceled' else 0)
le = LabelEncoder()
df['market segment type'] = le.fit_transform(df['market segment type'])

features = ['lead time', 'P-C', 'P-not-C', 'repeated', 'market segment type']
X = df[features].values
y = df['is_canceled'].values
X_train, X_test, y_train, y_test = train_test_split(X, y, test_size=0.2, random_state=42)

knn = KNeighborsClassifier(n_neighbors=5)
knn.fit(X_train, y_train)
y_pred = knn.predict(X_test)
accuracy = accuracy_score(y_test, y_pred)
report = classification_report(y_test, y_pred)

# ----------------------------
# Step 2: Create presentation
# ----------------------------
prs = Presentation()
blank_slide = prs.slide_layouts[5]

# Title slide
slide = prs.slides.add_slide(prs.slide_layouts[0])
slide.shapes.title.text = "Hotel Booking Cancellation Analysis"
slide.placeholders[1].text = "EDA + KNN Model (Auto-generated)"

# Helper to add slide with text
def add_text_slide(title, content):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    slide.placeholders[1].text = content

# Slide: Dataset Info
add_text_slide("Dataset Overview", f"Rows: {df.shape[0]}\nColumns: {df.shape[1]}\nNo missing values.\n\nTarget column: is_canceled")

# Slide: Model Info
add_text_slide("KNN Model Info", f"Features used: {', '.join(features)}\nk = 5\nDistance: Euclidean")

# Slide: Model Performance
add_text_slide("Model Performance", f"Accuracy: {accuracy:.2f}\n\nClassification Report:\n{report}")

# ----------------------------
# Step 3: Save plots and add them
# ----------------------------
def add_plot(title, plot_func, filename):
    plt.figure()
    plot_func()
    plt.title(title)
    plt.tight_layout()
    plt.savefig(filename)
    plt.close()
    
    slide = prs.slides.add_slide(blank_slide)
    slide.shapes.title.text = title
    slide.shapes.add_picture(filename, Inches(1), Inches(1.5), width=Inches(7))
    os.remove(filename)

# Plot functions
add_plot("Booking Cancellation Distribution", lambda: sns.countplot(x='is_canceled', data=df), "plot1.png")
add_plot("Lead Time vs Cancellation", lambda: sns.boxplot(x='is_canceled', y='lead time', data=df), "plot2.png")
add_plot("P-C vs Cancellation", lambda: sns.boxplot(x='is_canceled', y='P-C', data=df), "plot3.png")
add_plot("P-not-C vs Cancellation", lambda: sns.boxplot(x='is_canceled', y='P-not-C', data=df), "plot4.png")
add_plot("Repeated vs Cancellation", lambda: sns.countplot(x='repeated', hue='is_canceled', data=df), "plot5.png")
add_plot("Market Segment vs Cancellation", lambda: sns.countplot(x='market segment type', hue='is_canceled', data=df), "plot6.png")

# Confusion Matrix
cm = confusion_matrix(y_test, y_pred)
disp = ConfusionMatrixDisplay(confusion_matrix=cm, display_labels=["Not Canceled", "Canceled"])
disp.plot(cmap='Blues')
plt.title("Confusion Matrix")
plt.savefig("confusion.png")
plt.close()

slide = prs.slides.add_slide(blank_slide)
slide.shapes.title.text = "Confusion Matrix"
slide.shapes.add_picture("confusion.png", Inches(1), Inches(1.5), width=Inches(7))
os.remove("confusion.png")

# Save file
prs.save("Hotel_Cancellation_Analysis.pptx")
print("✅ PowerPoint saved as 'Hotel_Cancellation_Analysis.pptx'")
